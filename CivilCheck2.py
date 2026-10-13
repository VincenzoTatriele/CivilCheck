from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from tkinter import Tk, filedialog
import os

# 1. Scegliamo il computo da controllare
finestra = Tk()
finestra.withdraw()

percorso = filedialog.askopenfilename(
    title="Seleziona il computo metrico",
    filetypes=[("File Excel", "*.xlsx *.xlsm")]
)

finestra.destroy()

if not percorso:
    print("Operazione annullata.")
    raise SystemExit

# 2. Apriamo il file originale
file_excel = load_workbook(percorso, data_only=True)
foglio = file_excel.active

# 3. Creiamo il report
report = Workbook()

riepilogo = report.active
riepilogo.title = "Riepilogo"

anomalie_foglio = report.create_sheet("Anomalie")

statistiche_foglio = report.create_sheet("Statistiche")


riepilogo.append([
    "Codice",
    "Descrizione",
    "Quantità",
    "U.M.",
    "Prezzo unitario",
    "Importo",
    "Esito"
])

anomalie_foglio.append([
    "Riga originale",
    "Codice",
    "Problema"
])

statistiche_foglio.append([
    "Tipo di anomalia",
    "Numero di occorrenze"
])

totale_completo = 0
totale_calcolabile = 0
codici_visti = set()
conteggio_anomalie = {}
voci_complete = 0
voci_sospese = 0


# 4. Analizziamo ogni lavorazione
for riga in range(2, foglio.max_row + 1):

    codice = foglio.cell(riga, 1).value
    descrizione = foglio.cell(riga, 2).value
    quantita = foglio.cell(riga, 3).value
    prezzo = foglio.cell(riga, 4).value
    unita = foglio.cell(riga, 5).value

    valori = [codice, descrizione, quantita, prezzo, unita]

    # Ignoriamo le righe completamente vuote
    if all(v is None or str(v).strip() == "" for v in valori):
        continue

    problemi = []

    # Controlliamo i campi mancanti
    if codice is None or str(codice).strip() == "":
        problemi.append("Codice mancante")

    if descrizione is None or str(descrizione).strip() == "":
        problemi.append("Descrizione mancante")

    if quantita is None or quantita == "":
        problemi.append("Quantità mancante")

    if prezzo is None or prezzo == "":
        problemi.append("Prezzo unitario mancante")

    if unita is None or str(unita).strip() == "":
        problemi.append("Unità di misura mancante")

    # Controlliamo i codici duplicati
    if codice is not None and str(codice).strip() != "":
        codice_norm = str(codice).strip().upper()

        if codice_norm in codici_visti:
            problemi.append("Codice duplicato")
        else:
            codici_visti.add(codice_norm)
            

    # Verifichiamo quantità e prezzo
    numeri_validi = (
        isinstance(quantita, (int, float))
        and not isinstance(quantita, bool)
        and isinstance(prezzo, (int, float))
        and not isinstance(prezzo, bool)
    )

    importo = None

    if numeri_validi:
        if quantita < 0:
            problemi.append("Quantità negativa")

        if prezzo < 0:
            problemi.append("Prezzo negativo")

        importo = quantita * prezzo
        totale_calcolabile += importo

    elif quantita not in (None, "") and prezzo not in (None, ""):
        problemi.append("Quantità o prezzo non numerico")

    # Una voce è completa solo se non ha problemi
    if problemi:
        esito = "DA VERIFICARE"
        voci_sospese += 1
    else:
        esito = "COMPLETA"
        voci_complete += 1
        totale_completo += importo

    
    # Scriviamo la voce nel report
    riepilogo.append([
        codice,
        descrizione,
        quantita,
        unita,
        prezzo,
        importo,
        esito
    ])

    # Scriviamo tutte le anomalie e aggiorniamo i conteggi
    for problema in problemi:
        anomalie_foglio.append([
            riga,
            codice,
            problema
        ])

        if problema not in conteggio_anomalie:
            conteggio_anomalie[problema] = 0

        conteggio_anomalie[problema] += 1

# 5. Aggiungiamo il riepilogo economico
riga_totale = riepilogo.max_row + 2

riepilogo.cell(riga_totale, 4, "TOTALE VOCI COMPLETE")
riepilogo.cell(riga_totale, 5, totale_completo)

riepilogo.cell(riga_totale + 1, 4, "TOTALE CALCOLABILE")
riepilogo.cell(riga_totale + 1, 5, totale_calcolabile)


for problema, conteggio in sorted(conteggio_anomalie.items()):
    statistiche_foglio.append([
        problema,
        conteggio
    ])

# 6. Formattiamo i fogli
for foglio_report in [
    riepilogo,
    anomalie_foglio,
    statistiche_foglio
]:

    for cella in foglio_report[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="1F4E78")
        cella.alignment = Alignment(horizontal="center")

    foglio_report.freeze_panes = "A2"
    foglio_report.auto_filter.ref = (
        f"A1:{get_column_letter(foglio_report.max_column)}"
        f"{foglio_report.max_row}"
    )

    for colonna in foglio_report.columns:
        lettera = get_column_letter(colonna[0].column)
        larghezza = max(
            len(str(cella.value or ""))
            for cella in colonna
        )
        foglio_report.column_dimensions[lettera].width = min(
            larghezza + 3, 45
        )

# 7. Formattiamo gli importi
for riga in range(2, riepilogo.max_row + 1):
   riepilogo.cell(riga, 5).number_format = '#,##0.00'
   riepilogo.cell(riga, 6).number_format = '#,##0.00'



           # Evidenziamo le voci da verificare
for riga in range(2, riga_totale):

    if riepilogo.cell(riga, 7).value == "DA VERIFICARE":
        for cella in riepilogo[riga]:
            cella.fill = PatternFill(
                "solid",
                fgColor="FCE4D6"
            )

# 8. Salviamo il report accanto all'originale
cartella = os.path.dirname(percorso)
nome_file = os.path.splitext(os.path.basename(percorso))[0]

percorso_report = os.path.join(
    cartella,
    nome_file + "_report_civilcheck.xlsx"
)

report.save(percorso_report)

# 9. Mostriamo il risultato
print("\n--- RISULTATO CIVILCHECK ---")
print("Voci complete:", voci_complete)
print("Voci da verificare:", voci_sospese)
print("Totale voci complete:", round(totale_completo, 2), "euro")
print("Totale calcolabile:", round(totale_calcolabile, 2), "euro")
print("Report creato:", percorso_report)

      
